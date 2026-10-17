import os
import sys

def create_submission_word_doc(
    student_name: str = "Kumar Kshitij",
    registration_num: str = "230953488",
    github_link: str = "https://github.com/Kshiku12/Autoregressive-Text-Generation-RNN-LSTM-GRU",
    output_filename: str = "KumarKshitij_230953488_FGAI_Assignment2.docx"
):
    try:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.enum.table import WD_TABLE_ALIGNMENT
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
    except ImportError:
        print("python-docx not installed. Run 'pip install python-docx' first.")
        return

    doc = docx.Document()

    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    def set_cell_background(cell, hex_color):
        tcPr = cell._element.get_or_add_tcPr()
        shd = OxmlElement('w:shd')
        shd.set(qn('w:val'), 'clear')
        shd.set(qn('w:color'), 'auto')
        shd.set(qn('w:fill'), hex_color)
        tcPr.append(shd)

    # Document Header / Title
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Foundations of Generative AI (ICT4403)\nAssignment 2 Submission Document")
    title_run.font.name = "Calibri"
    title_run.font.size = Pt(20)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 78, 121)

    subtitle_p = doc.add_paragraph()
    subtitle_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_run = subtitle_p.add_run("Autoregressive Text Generation System using Simple RNN, Bidirectional RNN, LSTM, and GRU")
    sub_run.font.name = "Calibri"
    sub_run.font.size = Pt(13)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(89, 89, 89)

    doc.add_paragraph()

    # Step 3 Required Details Table
    details_heading = doc.add_heading("1. Submission Details (Required by Step 3)", level=2)
    details_heading.style.font.color.rgb = RGBColor(31, 78, 121)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    labels_data = [
        ("Student Name", student_name),
        ("Registration Number", registration_num),
        ("Assignment Title", "Foundations of Generative AI (ICT4403) - Assignment 2"),
        ("Title of Developed System", "Comparative Autoregressive Text Generation using Simple RNN, Bidirectional RNN, LSTM, and GRU Architectures"),
        ("Dataset Name and Link", "The Adventures of Sherlock Holmes by Arthur Conan Doyle\nhttps://www.gutenberg.org/ebooks/1661"),
        ("GitHub Repository Link", github_link)
    ]

    for i, (label, val) in enumerate(labels_data):
        row = table.rows[i]
        cell_lbl, cell_val = row.cells[0], row.cells[1]
        cell_lbl.width = Inches(2.2)
        cell_val.width = Inches(4.3)
        
        set_cell_background(cell_lbl, "F2F2F2")
        
        p_lbl = cell_lbl.paragraphs[0]
        r_lbl = p_lbl.add_run(label)
        r_lbl.font.bold = True
        r_lbl.font.size = Pt(11)
        r_lbl.font.name = "Calibri"

        p_val = cell_val.paragraphs[0]
        r_val = p_val.add_run(val)
        r_val.font.size = Pt(11)
        r_val.font.name = "Calibri"
        if "http" in val:
            r_val.font.color.rgb = RGBColor(5, 99, 193)
            r_val.font.underline = True

    doc.add_paragraph()

    # Executive Summary of Results
    h2 = doc.add_heading("2. Executive Summary & Experimental Results", level=2)
    h2.style.font.color.rgb = RGBColor(31, 78, 121)

    doc.add_paragraph(
        "This project develops and empirically evaluates an autoregressive next-word language modeling system across four fundamental recurrent neural architectures: "
        "Simple RNN, Bidirectional RNN, LSTM, and GRU. All four models were trained under strictly controlled, identical experimental conditions: "
        "same dataset split (85% train, 15% validation), vocabulary size (4,406 words, 97.14% token coverage), embedding dimension (128), hidden units (256), "
        "batch size (128), context window (20 words), learning rate (0.002 with Adam), and cross-entropy loss."
    )

    # Key Findings & Conclusion
    h3 = doc.add_heading("3. Key Findings & Best-Performing Model", level=2)
    h3.style.font.color.rgb = RGBColor(31, 78, 121)

    p_findings = doc.add_paragraph()
    p_findings.add_run("• Quantitative Superiority: ").bold = True
    p_findings.add_run("LSTM and GRU decisively outperform Simple RNN and Bidirectional RNN in validation loss and perplexity. LSTM achieved the lowest validation loss (4.8432) and lowest perplexity (126.88), with top-1 next-word accuracy of 18.79% and top-5 accuracy of 41.44%.\n")

    p_findings.add_run("• Architectural Efficiency: ").bold = True
    p_findings.add_run("GRU closely matches LSTM performance (Val Loss 4.8807, PPL 131.73, Top-1 Acc 18.12%) while requiring fewer parameters (1,992,758 vs. 2,091,574) and faster training time per epoch.\n")

    p_findings.add_run("• Temperature Influence: ").bold = True
    p_findings.add_run("At T = 0.2, models degenerate into repetitive loops. At T = 1.2, models produce high lexical diversity but lose grammatical coherence. Temperature T = 0.7 offers the optimal trade-off between coherence and diversity.\n")

    p_findings.add_run("• Final Selection: ").bold = True
    p_findings.add_run("LSTM is selected as the most effective model for autoregressive generation on this corpus.")

    doc.add_paragraph()

    # Instructions box for student
    instr_box = doc.add_table(rows=1, cols=1)
    instr_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = instr_box.rows[0].cells[0]
    c.width = Inches(6.5)
    set_cell_background(c, "EBF1F5")
    p_box = c.paragraphs[0]
    r_box_title = p_box.add_run("Submission Instructions for Student:\n")
    r_box_title.bold = True
    r_box_title.font.color.rgb = RGBColor(31, 78, 121)
    p_box.add_run(
        "1. Open this file and replace 'Student Name', 'Registration Number', and your GitHub Repository Link.\n"
        "2. Save this Word document with the exact name: YourName_RegNum_FGAI_Assignment2.docx (Step 3).\n"
        "3. Convert to PDF via File -> Save As -> PDF (Step 4).\n"
        "4. Upload the PDF file to the LMS assignment portal before the deadline (Step 5)."
    )

    doc.save(output_filename)
    print(f"Successfully generated Word submission document: {output_filename}")

def create_submission_pdf(
    student_name: str = "Kumar Kshitij",
    registration_num: str = "230953488",
    github_link: str = "https://github.com/Kshiku12/Autoregressive-Text-Generation-RNN-LSTM-GRU",
    output_filename: str = "KumarKshitij_230953488_FGAI_Assignment2.pdf"
):
    try:
        from reportlab.lib.pagesizes import letter
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle
        from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    except ImportError:
        print("reportlab not installed. Skipping direct PDF generation.")
        return

    doc = SimpleDocTemplate(
        output_filename,
        pagesize=letter,
        rightMargin=54,
        leftMargin=54,
        topMargin=54,
        bottomMargin=54,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        "TitleStyle",
        parent=styles["Heading1"],
        fontName="Helvetica-Bold",
        fontSize=18,
        leading=22,
        textColor=colors.HexColor("#1f4e79"),
        alignment=1,
    )

    sub_style = ParagraphStyle(
        "SubStyle",
        parent=styles["Normal"],
        fontName="Helvetica-Oblique",
        fontSize=11,
        leading=15,
        textColor=colors.HexColor("#595959"),
        alignment=1,
    )

    h2_style = ParagraphStyle(
        "H2Style",
        parent=styles["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=13,
        leading=17,
        textColor=colors.HexColor("#1f4e79"),
        spaceBefore=12,
        spaceAfter=6,
    )

    body_style = ParagraphStyle(
        "BodyStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=14,
        textColor=colors.HexColor("#262626"),
    )

    bold_label = ParagraphStyle(
        "BoldLabel",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#1f4e79"),
    )

    link_style = ParagraphStyle(
        "LinkStyle",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        textColor=colors.HexColor("#0563c1"),
    )

    story = []
    story.append(Paragraph("Foundations of Generative AI (ICT4403)<br/>Assignment 2 Submission Document", title_style))
    story.append(Spacer(1, 4))
    story.append(Paragraph("Autoregressive Text Generation System using Simple RNN, Bidirectional RNN, LSTM, and GRU", sub_style))
    story.append(Spacer(1, 14))

    story.append(Paragraph("1. Submission Details (Required by Step 3)", h2_style))

    table_data = [
        [Paragraph("Student Name", bold_label), Paragraph(student_name, body_style)],
        [Paragraph("Registration Number", bold_label), Paragraph(registration_num, body_style)],
        [Paragraph("Assignment Title", bold_label), Paragraph("Foundations of Generative AI (ICT4403) - Assignment 2", body_style)],
        [Paragraph("Title of Developed System", bold_label), Paragraph("Comparative Autoregressive Text Generation using Simple RNN, Bidirectional RNN, LSTM, and GRU Architectures", body_style)],
        [Paragraph("Dataset Name and Link", bold_label), Paragraph("The Adventures of Sherlock Holmes by Arthur Conan Doyle<br/><u>https://www.gutenberg.org/ebooks/1661</u>", link_style)],
        [Paragraph("GitHub Repository Link", bold_label), Paragraph(f"<u>{github_link}</u>", link_style)],
    ]

    t = Table(table_data, colWidths=[150, 350])
    t.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (0, -1), colors.HexColor("#f2f2f2")),
        ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#cccccc")),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(t)
    story.append(Spacer(1, 12))

    story.append(Paragraph("2. Executive Summary & Experimental Results", h2_style))
    story.append(Paragraph(
        "This project develops and empirically evaluates an autoregressive next-word language modeling system across four fundamental recurrent neural architectures: "
        "Simple RNN, Bidirectional RNN, LSTM, and GRU. All four models were trained under strictly controlled, identical experimental conditions: "
        "same dataset split (85% train, 15% validation), vocabulary size (4,406 words, 97.14% token coverage), embedding dimension (128), hidden units (256), "
        "batch size (128), context window (20 words), learning rate (0.002 with Adam), and cross-entropy loss.",
        body_style
    ))
    story.append(Spacer(1, 10))

    story.append(Paragraph("3. Key Findings & Best-Performing Model", h2_style))
    story.append(Paragraph(
        "<b>• Quantitative Superiority:</b> LSTM and GRU decisively outperform Simple RNN and Bidirectional RNN. LSTM achieved the lowest validation loss (4.8432) and lowest perplexity (126.88), with top-1 next-word accuracy of 18.79% and top-5 accuracy of 41.44%.<br/><br/>"
        "<b>• Architectural Efficiency:</b> GRU closely matches LSTM performance (Val Loss 4.8807, PPL 131.73, Top-1 Acc 18.12%) while requiring fewer parameters (1,992,758 vs. 2,091,574) and faster training time per epoch.<br/><br/>"
        "<b>• Temperature Influence:</b> At T = 0.2, models degenerate into repetitive loops. At T = 1.2, models produce high lexical diversity but lose grammatical coherence. Temperature T = 0.7 offers the optimal trade-off between coherence and diversity.<br/><br/>"
        "<b>• Final Selection:</b> LSTM is identified as the most effective model for autoregressive generation on this corpus.",
        body_style
    ))
    story.append(Spacer(1, 14))

    box_data = [[
        Paragraph(
            "<b>Submission Instructions for Student:</b><br/>"
            "1. Before uploading, ensure your Student Name, Registration Number, and GitHub Link are updated in the table above.<br/>"
            "2. Upload this final PDF file to the specified assignment submission section in the LMS before the deadline (Step 5).",
            body_style
        )
    ]]
    box_table = Table(box_data, colWidths=[500])
    box_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#ebf1f5")),
        ('BOX', (0, 0), (-1, -1), 1, colors.HexColor("#1f4e79")),
        ('TOPPADDING', (0, 0), (-1, -1), 8),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 8),
        ('LEFTPADDING', (0, 0), (-1, -1), 10),
        ('RIGHTPADDING', (0, 0), (-1, -1), 10),
    ]))
    story.append(box_table)

    doc.build(story)
    print(f"Successfully generated PDF submission document: {output_filename}")

if __name__ == "__main__":
    name = sys.argv[1] if len(sys.argv) > 1 else "Kumar Kshitij"
    reg = sys.argv[2] if len(sys.argv) > 2 else "230953488"
    gh = sys.argv[3] if len(sys.argv) > 3 else "https://github.com/Kshiku12/Autoregressive-Text-Generation-RNN-LSTM-GRU"
    create_submission_word_doc(name, reg, gh)
    # create_submission_pdf only if needed, do not overwrite user's custom pdf

